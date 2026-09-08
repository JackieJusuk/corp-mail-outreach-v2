# -*- coding: utf-8 -*-
"""
공개데이터 엑셀 파일을 파싱하여 SQLite DB(companies 테이블)에 적재하는 배치 스크립트.
requirements.md 2.3절(데이터 수집·저장 요구사항) 참조.

사용법:
    python scripts/import_excel.py --file <엑셀경로> [--db <DB경로>] \
        [--header-row 2] \
        [--mapping-overrides "원본헤더=표준필드명,원본헤더2=표준필드명2"] \
        [--letter-overrides "AM=email,C=business_reg_no"]

동작 순서:
    1. --header-row로 지정한 행(기본 1행)을 헤더로 읽고 표준 필드명으로 정규화한다
       (column_mapping.py). 사업자등록번호/지역/상호/대표자명/이메일 5개 필드만 신경 쓰며,
       그 외 컬럼(재무정보, 주주정보 등 무관한 컬럼)은 매핑 안 돼도 그냥 무시한다.
       일부 공개데이터/구매 DB 엑셀은 1행이 병합된 대분류 제목이고 실제 항목명은
       2행에 있는 경우가 있다 — 이때는 --header-row 2 를 지정한다.
    2. **필수 필드**(사업자등록번호, 상호)가 안 보이면 적재를 중단한다.
       -> 로컬 Claude Code가 안내를 보고 사용자에게 해당 컬럼이 무엇인지 물어본 뒤,
          --mapping-overrides 또는 --letter-overrides 옵션으로 다시 실행한다.
       **선택 필드**(지역, 대표자명, 이메일)가 안 보이면 중단하지 않고 안내만 한다.
    3. 사업자등록번호(Primary Key)가 이미 DB에 있으면 해당 레코드는 건너뛴다 (skip, 갱신 없음).
    4. 이메일이 비어 있으면 email_status='missing'으로 저장한다 (발송 대상에서 자동 제외됨).

참고: 매핑 옵션은 JSON이 아니라 "key=value,key2=value2" 형식이다. Windows
PowerShell/cmd는 큰따옴표가 포함된 JSON을 명령줄 인자로 넘기면 셸마다 따옴표
처리 방식이 달라 깨지기 쉬워서, 따옴표가 필요 없는 이 형식을 쓴다.
"""
import argparse
import sqlite3
import sys
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
from column_mapping import (  # noqa: E402
    OPTIONAL_FIELDS,
    REQUIRED_FIELDS,
    parse_kv_pairs,
    resolve_columns,
)

DEFAULT_DB_PATH = str(Path(__file__).resolve().parent.parent / "corp_mail_outreach.db")


def parse_args():
    parser = argparse.ArgumentParser(description="공개데이터 엑셀 파일을 DB에 적재")
    parser.add_argument("--file", required=True, help="적재할 엑셀 파일 경로")
    parser.add_argument("--db", default=DEFAULT_DB_PATH, help="SQLite DB 파일 경로")
    parser.add_argument(
        "--header-row",
        type=int,
        default=1,
        help="실제 헤더(항목명)가 있는 행 번호, 1부터 시작 (기본 1). "
        "1행이 병합된 대분류 제목이고 2행에 실제 항목명이 있으면 2로 지정",
    )
    parser.add_argument(
        "--mapping-overrides",
        default="",
        help='매핑 실패 헤더에 대한 수동 매핑 (헤더 텍스트 기준), 예: "업체소재지=region,담당자메일=email"',
    )
    parser.add_argument(
        "--letter-overrides",
        default="",
        help='엑셀 열 문자 기준 수동 매핑, 예: "AM=email,C=business_reg_no" (헤더 텍스트가 특이할 때 사용, mapping-overrides보다 우선)',
    )
    return parser.parse_args()


def load_rows(file_path, header_row=1):
    """엑셀 파일을 읽어 (헤더 목록, 데이터 행 목록)을 반환한다. header_row는 1부터 시작."""
    wb = load_workbook(filename=file_path, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < header_row:
        return [], []
    headers = [str(h).strip() if h is not None else "" for h in rows[header_row - 1]]
    data_rows = rows[header_row:]
    return headers, data_rows


def clean(value):
    """엑셀 셀 값을 문자열로 정리한다 (앞뒤 공백 제거, 빈 값은 None)."""
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def main():
    args = parse_args()
    overrides = parse_kv_pairs(args.mapping_overrides)
    letter_overrides = parse_kv_pairs(args.letter_overrides)

    headers, data_rows = load_rows(args.file, args.header_row)
    if not headers:
        print("엑셀 파일에 데이터가 없습니다 (--header-row 값이 실제 파일보다 큰 건 아닌지 확인하세요).")
        sys.exit(1)

    mapping = resolve_columns(headers, overrides, letter_overrides)
    mapped_fields = set(mapping.values())

    missing_required = [f for f in REQUIRED_FIELDS if f not in mapped_fields]
    if missing_required:
        print(f"필수 필드를 찾지 못했습니다: {missing_required}")
        print("--mapping-overrides(헤더텍스트) 또는 --letter-overrides(열문자)로 해당 필드의 컬럼을 지정해 다시 실행하세요.")
        print(f"참고로 이번에 읽은 헤더 행({args.header_row}행) 전체 목록:")
        for idx, header in enumerate(headers):
            print(f"  [{idx + 1}열] {header!r}")
        sys.exit(2)

    missing_optional = [f for f in OPTIONAL_FIELDS if f not in mapped_fields]
    if missing_optional:
        print(f"참고: 다음 선택 필드는 식별되지 않아 비어서 저장됩니다: {missing_optional}")
        print("필요하면 --mapping-overrides 또는 --letter-overrides로 지정할 수 있습니다. (적재는 계속 진행됩니다)")

    conn = sqlite3.connect(args.db)
    cur = conn.cursor()

    inserted, skipped_dup, skipped_no_key = 0, 0, 0

    for row in data_rows:
        record = {}
        for idx, field in mapping.items():
            record[field] = row[idx] if idx < len(row) else None

        business_reg_no = clean(record.get("business_reg_no"))
        if not business_reg_no:
            skipped_no_key += 1
            continue

        email = clean(record.get("email"))
        email_status = "present" if email else "missing"

        cur.execute(
            """
            INSERT OR IGNORE INTO companies
                (business_reg_no, region, company_name, representative_name, email, email_status, source_file)
            VALUES (?, ?, ?, ?, ?, ?, ?)
            """,
            (
                business_reg_no,
                clean(record.get("region")),
                clean(record.get("company_name")),
                clean(record.get("representative_name")),
                email,
                email_status,
                Path(args.file).name,
            ),
        )
        if cur.rowcount == 1:
            inserted += 1
        else:
            skipped_dup += 1

    conn.commit()
    conn.close()

    print(
        f"적재 완료: 신규 {inserted}건, 중복(기존 사업자번호) 스킵 {skipped_dup}건, "
        f"사업자등록번호 없음 스킵 {skipped_no_key}건"
    )


if __name__ == "__main__":
    main()
